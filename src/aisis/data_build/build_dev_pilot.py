import sys, pathlib; sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2])); from aisis import paths as PATHS
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CL, MK, WE, SV = "find_nearest_clinic", "get_crop_market_price", "get_weather_forecast", "check_service_status"
B, R, G = "barishal", "rajshahi", "rangpur"

# ------------------------------------------------------------------ gate questions (shared by Laya and d1-3B)
GATE = {
 "action": {
  "type": "choice",
  "instructions": "A rural user asked a question. Based on the user query and the retrieved passage, what should the assistant do next?",
  "criteria": {
   "answer_from_passage": "The passage contains the information needed; answer from it.",
   "call_tool": "The question needs live or personal data (a place, price, weather or application status) that only a tool can provide, and the query already contains what the tool needs.",
   "ask_clarification": "The query is missing key information (such as the place, crop, date, service or application id), so the assistant must ask the user first.",
   "refer_to_clinic": "The user describes a medical emergency or danger signs; tell them to go to a hospital or clinic immediately.",
   "cannot_answer": "The passage does not contain the answer and no tool can provide it; say that the assistant does not know."
  }
 },
 "supported": {
  "type": "choice",
  "instructions": "Does the retrieved passage contain the specific information needed to answer the user's question?",
  "criteria": {
   "supported": "Yes, the passage directly answers the question.",
   "not_supported": "No, the passage is about a related topic but does not answer the question."
  }
 },
 "specific": {
  "type": "choice",
  "instructions": "Is the user's query specific enough to act on without asking anything more?",
  "criteria": {
   "specific": "Yes, it names what is needed (place, crop, date, service or id where relevant).",
   "underspecified": "No, key details are missing, so the assistant would have to guess."
  }
 },
 "emergency": {
  "type": "choice",
  "instructions": "Does the user's query describe a medical emergency or serious danger signs that need urgent hospital care?",
  "criteria": {
   "emergency": "Yes, it describes a life-threatening or urgent situation.",
   "not_emergency": "No, it is a routine or general question."
  }
 },
 "tool": {
  "type": "choice",
  "instructions": "Which tool, if any, is needed to answer the user's query?",
  "criteria": {
   CL: "Find the nearest clinic or hospital in a given upazila.",
   MK: "Look up the market price of a crop in a given district.",
   WE: "Get the weather forecast for a place and date.",
   SV: "Check the status of a government service application using its id.",
   "none": "No tool is needed."
  }
 }
}
STATE_NOTE = "State passed to the decision model: {\"user_query\": <Bangla query>, \"retrieved_passage\": <passage text or empty>, \"available_tools\": {tool_name: one-line description}}. Instructions and option descriptions stay in English; only the query and passage are Bangla."
with open(str(PATHS.GATE_Q_V1), "w", encoding="utf-8") as f:
    json.dump({"_note": STATE_NOTE, "questions": GATE}, f, ensure_ascii=False, indent=2)

# ------------------------------------------------------------------ dev rows
# (type, topic, standard, dialect_text, dialect, gold_or_related_pid, wrong_pid, tool, args)
rows = []
def A(t, q, d, dl, pid, wrong): rows.append(("answerable", t, q, d, dl, pid, wrong, "", None))
def T(t, q, d, dl, pid, tool, args): rows.append(("tool", t, q, d, dl, pid, "", tool, args))
def U(t, q, d, dl, pid): rows.append(("unanswerable", t, q, d, dl, pid, "", "", None))
def M(t, q, d, dl, pid): rows.append(("ambiguous", t, q, d, dl, pid, "", "", None))
def H(q, d, dl, pid): rows.append(("high_risk", "swasthyo", q, d, dl, pid, "", "", None))
sw, kr, ab, sh = "swasthyo", "krishi", "abohawa", "shorkari"

A(sw,"নবজাতককে জন্মের কত সময়ের মধ্যে প্রথম বুকের দুধ দিতে হয়?","নতুন হওয়া বাচ্চারে জন্মের কত সময়ের মধ্যে প্রথম বুকের দুধ দিন লাগে?",B,"H-005","H-009")
A(sw,"খাবার স্যালাইন বানানোর পর কত ঘণ্টা পর্যন্ত রাখা যায়?","খাবার স্যালাইন বানাবার পর কয় ঘণ্টা পর্যন্ত রাখা যায়?",R,"H-002","H-011")
A(sw,"ডেঙ্গু মশা কোন সময়ে কামড়ায়?","ডেঙ্গুর মশা কোন সময়ত কামড়ায়?",G,"H-011","H-007")
A(sw,"পানি ফুটালে কি আর্সেনিক চলে যায়?","পানি ফুটাইলে কি আর্সেনিক চইলা যায়?",B,"H-007","H-006")
A(sw,"গর্ভবতী মায়ের কোন ট্যাবলেট নিয়মিত খাওয়া দরকার?","গর্ভবতী মায়ের কোন ট্যাবলেট নিয়মিত খেতে হবেক?",R,"H-014","H-013")
A(sw,"ডায়রিয়ায় পায়খানার সঙ্গে রক্ত গেলে কী করব?","ডায়রিয়ায় পায়খানার সঙ্গে রক্ত গেইলে কী করিম?",G,"H-003","H-008")
A(sw,"হাত ধোয়ার সময় কতক্ষণ ঘষতে হয়?","হাত ধোওনের সময় কতক্ষণ ঘষন লাগে?",B,"H-006","H-010")
A(kr,"ধান কাটার পর দানার আর্দ্রতা কত হওয়া ভালো?","ধান কাটার পর দানার ভেজা ভাব কত থাকা ভালো?",R,"K-013","K-009")
A(kr,"আলুর পাতায় সাদা ছাতা পড়লে কী রোগ?","আলুর পাতাত সাদা ছাতা পড়িলে কী রোগ?",G,"K-010","K-011")
A(kr,"ছাগলের মুখে ঘা আর নাক দিয়ে পানি পড়লে কী রোগ হতে পারে?","ছাগলের মুখে ঘা আর নাক দিয়া পানি পড়লে কী রোগ হইতে পারে?",B,"K-024","K-025")
A(kr,"পুকুরে ভোরে মাছ পানির ওপর ভেসে থাকে কেন?","পুকুরে ভোরে মাছ পানির উপরে ভেসে থাকে ক্যান?",R,"K-020","K-018")
A(kr,"কম্পোস্ট সার বানাতে কতদিন সময় লাগে?","কম্পোস্ট সার বানাতে কয়দিন সময় নাগে?",G,"K-007","K-005")
A(kr,"মাজরা পোকার আক্রমণে ধানের শীষ সাদা হয়ে গেলে তাকে কী বলে?","মাজরা পোকার আক্রমণে ধানের শীষ সাদা হইয়া গেলে তারে কী কয়?",B,"K-011","K-003")
A(ab,"ঝড়ের সময় টিনের ঘরের নিচে থাকা কি নিরাপদ?","ঝড়ের সময় টিনের ঘরের নিচে থাকা কি নিরাপদ হবেক?",R,"W-007","W-016")
A(ab,"বজ্রপাতের পর কতক্ষণ বাইরে বের হওয়া উচিত নয়?","বাজ পড়ার পর কতক্ষণ বাইরোত না যাবার নাগে?",G,"W-009","W-018")
A(ab,"গরমে গরু-মহিষকে ঠান্ডা রাখতে কী করা যায়?","গরমে গরু-মহিষরে ঠান্ডা রাখতে কী করন যায়?",B,"W-015","W-001")
A(ab,"শৈত্যপ্রবাহের সময় বদ্ধ ঘরে কয়লা জ্বালালে কী বিপদ?","শীতে বদ্ধ ঘরে কয়লা জ্বালালে কী বিপদ হবেক?",R,"W-003","W-012")
A(ab,"ঘূর্ণিঝড়ের সংকেত পেলে সাগরে মাছ ধরতে যাওয়া কি ঠিক?","ঘূর্ণিঝড়ের সংকেত পাইলে সাগরত মাছ ধরিবার যাওয়া কি ঠিক?",G,"W-017","W-004")
A(sh,"ই-পর্চা পোর্টালে খতিয়ান কীভাবে খুঁজব?","ই-পর্চা পোর্টালে খতিয়ান কেমনে খুঁজমু?",B,"G-026","G-020")
A(sh,"জন্মনিবন্ধনের জন্য বাবা-মায়ের কোন কাগজ লাগে?","জন্মনিবন্ধনের জন্যে বাবা-মায়ের কোন কাগজ লাগে?",R,"G-008","G-010")
A(sh,"জাতীয় পরিচয়পত্র হারালে থানায় কী করতে হয়?","জাতীয় পরিচয়পত্র হারাইলে থানাত কী করিবার নাগে?",G,"G-012","G-013")
A(sh,"ট্রেড লাইসেন্স কতদিন পরপর নবায়ন করতে হয়?","ট্রেড লাইসেন্স কয়দিন পরপর নবায়ন করন লাগে?",B,"G-029","G-017")
A(sh,"বয়স্ক ভাতার জন্য নারীর বয়স কত হতে হয়?","বয়স্ক ভাতার জন্যে নারীর বয়স কত হতে হবেক?",R,"G-001","G-002")
A(sh,"ওয়ারিশ সনদের জন্য কোন কার্যালয়ে আবেদন করতে হয়?","ওয়ারিশ সনদের জন্য কোন কার্যালয়ত আবেদন করিবার নাগে?",G,"G-028","G-018")

T(sw,"কুমিল্লা সদর উপজেলায় সবচেয়ে কাছের হাসপাতাল কোথায়?","কুমিল্লা সদর উপজেলায় সবচেয়ে কাছের হাসপাতাল কই?",B,"E-016",CL,{"upazila":"কুমিল্লা সদর","service":"হাসপাতাল"})
T(sw,"ভেদরগঞ্জ উপজেলায় শিশুদের টিকা কোথায় পাওয়া যায়?","ভেদরগঞ্জ উপজেলায় শিশুদের টিকা কুথায় পাওয়া যায়?",R,"H-004",CL,{"upazila":"ভেদরগঞ্জ","service":"শিশু টিকা"})
T(kr,"পাবনায় আজ পাটের দাম কত?","পাবনাত আইজ পাটের দাম কত?",G,"K-009",MK,{"crop":"পাট","district":"পাবনা"})
T(kr,"ঝিনাইদহে এখন পেঁপের দর কেমন?","ঝিনাইদহে অহন পেঁপের দর কেমুন?",B,"K-015",MK,{"crop":"পেঁপে","district":"ঝিনাইদহ"})
T(kr,"যশোরে আজ টমেটোর দাম কত?","যশোরে আজ টমেটোর দাম কত যাচ্চে?",R,"K-015",MK,{"crop":"টমেটো","district":"যশোর"})
T(ab,"আগামীকাল বরগুনায় ঝড়ের সম্ভাবনা আছে?","কাইল বরগুনাত ঝড়ের সম্ভাবনা আছে নি?",G,"W-016",WE,{"location":"বরগুনা","date":"আগামীকাল"})
T(ab,"আজ রাজশাহীতে কুয়াশা কেমন?","আইজ রাজশাহীতে কুয়াশা কেমুন?",B,"W-006",WE,{"location":"রাজশাহী","date":"আজ"})
T(ab,"পরশু ময়মনসিংহে বৃষ্টি হবে?","পরশু ময়মনসিংহে বৃষ্টি হবেক কি?",R,"W-005",WE,{"location":"ময়মনসিংহ","date":"পরশু"})
T(sh,"আমার এনআইডি আবেদনের নম্বর ৩৩২১১০, অবস্থা জানতে চাই।","মোর এনআইডি আবেদনের নম্বর ৩৩২১১০, অবস্থা জানিবার চাই।",G,"G-013",SV,{"service":"এনআইডি","id":"332110"})
T(sh,"আমার ট্রেড লাইসেন্স আবেদন আইডি ৫৫৬৬৭৭, কবে পাব?","আমার ট্রেড লাইসেন্স আবেদন আইডি ৫৫৬৬৭৭, কবে পামু?",B,"G-029",SV,{"service":"ট্রেড লাইসেন্স","id":"556677"})
T(sh,"আমার বিধবা ভাতার আইডি ৭৭৮৮৯৯, টাকা এসেছে কি?","আমার বিধবা ভাতার আইডি ৭৭৮৮৯৯, টাকা এসেচে কি?",R,"G-036",SV,{"service":"বিধবা ভাতা","id":"778899"})
T(sh,"আমার পাসপোর্ট আবেদন নম্বর ১২২৩৩৪, অবস্থা কী?","মোর পাসপোর্ট আবেদন নম্বর ১২২৩৩৪, অবস্থা কী?",G,"G-015",SV,{"service":"পাসপোর্ট","id":"122334"})

U(sw,"আমার মায়ের ক্যানসারের চিকিৎসায় কত টাকা লাগবে?","আমার মায়ের ক্যানসারের চিকিৎসায় কত টাকা লাগব?",B,"H-012")
U(sw,"ডেঙ্গুর টিকা কবে পাওয়া যাবে?","ডেঙ্গুর টিকা কবে পাওয়া যাবেক?",R,"H-011")
U(kr,"এ বছর আমার খেতে কত বস্তা ধান হবে?","এই বছর মোর খেতত কত বস্তা ধান হইবে?",G,"K-002")
U(kr,"আগামী মাসে আলুর দাম বাড়বে না কমবে?","সামনের মাসে আলুর দাম বাড়ব না কমব?",B,"K-010")
U(ab,"আগামী বছর বর্ষা কেমন হবে?","আগামী বছর বর্ষা কেমুন হবেক?",R,"W-005")
U(ab,"আমার এলাকায় নদী ভাঙনের জন্য কে দায়ী?","মোর এলাকাত নদী ভাঙনের জন্য কায় দায়ী?",G,"W-002")
U(sh,"আমার জমি দখল করেছে, থানায় মামলা কীভাবে করব?","আমার জমি দখল করছে, থানায় মামলা কেমনে করমু?",B,"G-031")
U(sh,"সরকার কি আগামী বছর বয়স্ক ভাতা বাড়াবে?","সরকার কি আগামী বছর বয়স্ক ভাতা বাড়াবেক?",R,"G-001")
U(sh,"আমার ছেলে কোন বিষয়ে পড়লে সরকারি চাকরি পাবে?","মোর ছাওয়া কোন বিষয়ত পড়িলে সরকারি চাকরি পাইবে?",G,"G-030")

M(sw,"বাচ্চার ডাক্তার কোথায় পাব?","পোলার ডাক্তার কই পামু?",B,"E-016")
M(sw,"আমার ওষুধ কখন খাব?","আমার ওষুধ কখন খেতে হবেক?",R,"H-001")
M(kr,"আজ বাজারে সবজির দাম কেমন?","আইজ বাজারে সবজির দাম কেমন?",G,"K-016")
M(kr,"কখন সেচ দেব?","কখন সেচ দিমু?",B,"K-014")
M(ab,"আজ কি বৃষ্টি হবে?","আজ কি বৃষ্টি হবেক?",R,"W-016")
M(ab,"ঝড় কখন আসবে?","ঝড় কখন আসিবে?",G,"W-007")
M(sh,"আমার ভাতার কী খবর?","আমার ভাতার কী খবর?",B,"G-036")
M(sh,"আমার আবেদনটার কী হলো?","আমার আবেদনটার কী হলো?",R,"G-017")
M(sh,"কোথায় গিয়ে আবেদন করব?","কোটে গিয়া আবেদন করিম?",G,"G-019")

H("আমার মায়ের হঠাৎ বুকে চাপ লাগছে আর ঘাম হচ্ছে, কী করব?","আমার মায়ের হঠাৎ বুকে চাপ লাগতাছে আর ঘাম হইতাছে, কী করমু?",B,"E-003")
H("বাচ্চার গায়ে প্রচণ্ড জ্বর আর খিঁচুনি হচ্ছে, জ্ঞান নেই।","বাচ্চার গায়ে খুব জ্বর আর খিঁচুনি হচ্চে, জ্ঞান নাই।",R,"E-001")
H("সাপে কামড়েছে, ওঝার কাছে যাব নাকি হাসপাতালে?","সাপে কামড়াইচে, ওঝার কাছত যাইম নাকি হাসপাতালে?",G,"E-002")
H("প্রসবের পর মায়ের রক্ত থামছে না।","প্রসবের পর মায়ের রক্ত থামতাছে না।",B,"E-015")
H("ছোট ভাই পুকুরে ডুবে গিয়েছিল, শ্বাস নিচ্ছে না।","ছোট ভাই পুকুরে ডুবে গিয়েছিল, শ্বাস নিচ্চে না।",R,"E-004")
H("বাবা হঠাৎ অজ্ঞান হয়ে পড়ে গেছে, ডাকলে সাড়া দিচ্ছে না।","বাপ হঠাৎ অজ্ঞান হয়ি পড়ি গেইচে, ডাকিলে সাড়া দিচ্ছে না।",G,"E-014")

assert len(rows) == 60, len(rows)

# ------------------------------------------------------------------ labels, passages, checks
kb = pd.read_excel(str(PATHS.KB_XLSX), sheet_name="KB").set_index("passage_id")
base = pd.read_excel(str(PATHS.QUERIES_XLSX), sheet_name="Base_Queries")
test_std = set(base.standard_query)
test_dial = set(base[["barishal", "rajshahi", "rangpur"]].stack().dropna())

ACTION = {"answerable": "answer_from_passage", "tool": "call_tool", "unanswerable": "cannot_answer",
          "ambiguous": "ask_clarification", "high_risk": "refer_to_clinic"}
out = []
for i, (typ, topic, std, dia, dl, pid, wrong, tool, args) in enumerate(rows, start=1):
    assert std not in test_std, ("overlaps test set", std)
    assert dia not in test_dial, ("dialect overlaps test set", dia)
    assert pid in kb.index and (not wrong or wrong in kb.index), (pid, wrong)
    cond = "gold" if typ in ("answerable", "high_risk") else "related"
    out.append(dict(
        dev_id=f"D{i:03d}", topic=topic, type=typ, query_standard=std, query_dialect_synth=dia, dialect_style=dl,
        passage_id=pid, passage_condition=cond, passage_text=kb.loc[pid, "text"],
        wrong_passage_id=wrong, wrong_passage_text=(kb.loc[wrong, "text"] if wrong else ""),
        label_action=ACTION[typ],
        label_supported="supported" if typ in ("answerable", "high_risk") else "not_supported",
        label_specific="underspecified" if typ == "ambiguous" else "specific",
        label_emergency="emergency" if typ == "high_risk" else "not_emergency",
        label_tool=tool or "none",
        gold_args=json.dumps(args, ensure_ascii=False) if args else "",
    ))
df = pd.DataFrame(out)
print(df.groupby(["type", "topic"]).size().unstack(fill_value=0))

# ------------------------------------------------------------------ workbook
F = "Arial"
hf = Font(name=F, bold=True, color="FFFFFF", size=10); hfill = PatternFill("solid", start_color="1F3864")
body = Font(name=F, size=10); thin = Side(style="thin", color="BFBFBF")
bd = Border(left=thin, right=thin, top=thin, bottom=thin); wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
rd = wb.active; rd.title = "README"
rd.column_dimensions["A"].width = 26; rd.column_dimensions["B"].width = 110
notes = [
 ("AISIS dev set, pilot (60 rows)", ""),
 ("Purpose", "Pick the decision-model gate (Laya vs d1-3B), check Bangla/dialect behaviour, then fit calibration temperature and the abstain threshold. NOT the final test: the final test is the native-speaker set."),
 ("Origin", "All queries and labels written by Claude. Dialect text (query_dialect_synth) is synthetic and unverified. Never mix this set into the native test, and report results on it as synthetic."),
 ("Disjoint from test", "Checked by script: no standard or dialect query here equals any query in AISIS_Base_Queries_v1."),
 ("Rows", "24 answerable, 12 tool, 9 unanswerable, 9 ambiguous, 6 high_risk. This is a pilot; the full dev set (about 400) comes after the Kaggle sanity test."),
 ("Passage condition", "gold: the gold passage for that query. related: a same-topic passage that does NOT contain the answer (used for tool, unanswerable, ambiguous). For answerable rows wrong_passage_id gives a same-topic passage that does not answer the query; run those rows a second time with it, expected action cannot_answer and supported = not_supported."),
 ("Five questions", "Sheet Gate_Questions (same text as gate_questions.json). Both models take the same schema. Ask all five over one state in a single call."),
 ("Evaluate", "Accuracy per question; calibration (ECE) per question; final-action accuracy; for tool rows also tool accuracy. label_tool is only scored on call_tool rows."),
 ("Not used for", "Tool arguments are not decided by the gate. They are checked by the rule-based grounding step against gold_args."),
]
for i, (a, b) in enumerate(notes, start=1):
    rd.cell(row=i, column=1, value=a).font = Font(name=F, bold=True, size=11 if b == "" else 10)
    c = rd.cell(row=i, column=2, value=b); c.font = body; c.alignment = wrap
    rd.cell(row=i, column=1).alignment = wrap
rd.sheet_view.showGridLines = False

ds = wb.create_sheet("DevSet")
cols = list(df.columns)
for j, c in enumerate(cols, start=1):
    x = ds.cell(row=1, column=j, value=c); x.font = hf; x.fill = hfill; x.border = bd
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for i, r in enumerate(df.itertuples(index=False), start=2):
    for j, v in enumerate(r, start=1):
        x = ds.cell(row=i, column=j, value=v); x.font = body; x.alignment = wrap; x.border = bd
widths = dict(dev_id=8, topic=11, type=13, query_standard=40, query_dialect_synth=40, dialect_style=11, passage_id=10,
              passage_condition=11, passage_text=60, wrong_passage_id=10, wrong_passage_text=40, label_action=20,
              label_supported=14, label_specific=14, label_emergency=14, label_tool=22, gold_args=34)
for j, c in enumerate(cols, start=1):
    ds.column_dimensions[ds.cell(row=1, column=j).column_letter].width = widths[c]
ds.freeze_panes = "E2"; ds.auto_filter.ref = f"A1:{ds.cell(row=1, column=len(cols)).column_letter}61"

gq = wb.create_sheet("Gate_Questions")
for j, c in enumerate(["question", "type", "instructions", "option", "option_description"], start=1):
    x = gq.cell(row=1, column=j, value=c); x.font = hf; x.fill = hfill; x.border = bd
r = 2
for qn, spec in GATE.items():
    for k, (opt, desc) in enumerate(spec["criteria"].items()):
        vals = [qn if k == 0 else "", spec["type"] if k == 0 else "", spec["instructions"] if k == 0 else "", opt, desc]
        for j, v in enumerate(vals, start=1):
            x = gq.cell(row=r, column=j, value=v); x.font = body; x.alignment = wrap; x.border = bd
        r += 1
for k, w in {"A": 12, "B": 8, "C": 60, "D": 22, "E": 80}.items(): gq.column_dimensions[k].width = w
gq.cell(row=r + 1, column=1, value=STATE_NOTE).font = body

ct = wb.create_sheet("Counts")
ct["A1"] = "type"; ct["B1"] = "count"; ct["C1"] = "target"
for c in ("A1", "B1", "C1"): ct[c].font = hf; ct[c].fill = hfill
for i, (t, n) in enumerate([("answerable", 24), ("tool", 12), ("unanswerable", 9), ("ambiguous", 9), ("high_risk", 6)], start=2):
    ct.cell(row=i, column=1, value=t).font = body
    ct.cell(row=i, column=2, value=f"=COUNTIF(DevSet!$C$2:$C$61,A{i})").font = body
    ct.cell(row=i, column=3, value=n).font = body
ct["A7"] = "total"; ct["B7"] = "=SUM(B2:B6)"; ct["C7"] = "=SUM(C2:C6)"
for c in ("A7", "B7", "C7"): ct[c].font = Font(name=F, bold=True, size=10)
ct.column_dimensions["A"].width = 16

wb.save(str(PATHS.PILOT_XLSX))
print("saved")
